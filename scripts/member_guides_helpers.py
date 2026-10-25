import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_banner(doc, member_name, student_id, role, presentation_order, target_time):
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ ĐỒNG NAI (DNTU)\nKHOA CÔNG NGHỆ THÔNG TIN – HỌC PHẦN WEB API (TH71108)")
    r_uni.bold = True
    r_uni.font.size = Pt(10)
    r_uni.font.color.rgb = RGBColor(34, 100, 172)
    p_uni.paragraph_format.space_after = Pt(4)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"CẨM NANG LUYỆN NÓI & PHẢN BIỆN CÁ NHÂN\n{member_name.upper()} ({student_id})")
    r_title.bold = True
    r_title.font.size = Pt(15)
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    p_title.paragraph_format.space_after = Pt(4)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(f"Vai trò: {role} | Thứ tự thuyết trình: {presentation_order} | Thời lượng mục tiêu: {target_time}\nDự án: Hệ thống Quản lý Hoạt động Hội nhập – VISIONARY")
    r_meta.italic = True
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    p_meta.paragraph_format.space_after = Pt(14)

def add_callout(doc, text, title="LƯU Ý QUAN TRỌNG", border_color="2264AC", bg_color="F0F6FC"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(34, 100, 172)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_code_box(doc, code_str, caption=""):
    if caption:
        p_cap = doc.add_paragraph()
        r_cap = p_cap.add_run(f"💻 Mã nguồn minh chứng: {caption}")
        r_cap.bold = True
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = RGBColor(34, 100, 172)
        p_cap.paragraph_format.space_after = Pt(2)
        
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_str)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(226, 232, 240)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Helper functions ready.")
